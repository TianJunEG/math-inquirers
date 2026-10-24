#!/usr/bin/env python3
"""Build the operations-kit PDFs and Word files into ops/out/.

  parent_add_to_calendar.pdf            one page for parents
  makeup_note_parents.pdf               one page for parents
  end_of_program_report_L1..L6.pdf/.docx  one page per level (topics pre-filled), printable + editable
  tutor_notes_sheet.pdf                 PRIVATE tutor notes (feeds the report)
  tutor_observation_sheet.pdf           A4 timing/observation sheet for dry runs
Run: python3 ops/make_docs.py
"""
import json
from pathlib import Path
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                                KeepTogether)

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "ops" / "out"
OUT.mkdir(parents=True, exist_ok=True)
LL = json.loads((ROOT / "website/data/level_lessons.json").read_text())
SCHED = json.loads((ROOT / "website/data/beta_schedule.json").read_text())

INK = colors.HexColor("#1f2a44")
ACC = colors.HexColor("#2b6cb0")
GREY = colors.HexColor("#555555")
LINE = colors.HexColor("#999999")
WA = "+65 9880 5962"
CO = "Math Inquirers, by Preciouslyours Pte Ltd, Singapore"

H1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=INK, spaceAfter=3)
H2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=ACC, spaceBefore=7, spaceAfter=2)
B = ParagraphStyle("b", fontName="Helvetica", fontSize=9.6, leading=13, textColor=INK)
BS = ParagraphStyle("bs", parent=B, fontSize=8.4, leading=11, textColor=GREY)
CELL = ParagraphStyle("c", parent=B, fontSize=8.5, leading=10.5)
CELLB = ParagraphStyle("cb", parent=CELL, fontName="Helvetica-Bold")
BUL = ParagraphStyle("bul", parent=B, leftIndent=11, bulletIndent=1, spaceAfter=1.5)
BOX = "[   ]"


def P(t, s=B):
    return Paragraph(t, s)


def bullets(items, style=BUL):
    return [Paragraph(i, style, bulletText="•") for i in items]


def foot(c, d):
    c.saveState()
    c.setFont("Helvetica", 7.5); c.setFillColor(GREY)
    w = c._pagesize[0]
    c.drawString(14 * mm, 8 * mm, CO)
    c.drawRightString(w - 14 * mm, 8 * mm, getattr(d, "_foot", ""))
    c.restoreState()


def build(path, story, size=A4, margin=14, footer="", title=""):
    d = SimpleDocTemplate(str(path), pagesize=size, leftMargin=margin * mm, rightMargin=margin * mm,
                          topMargin=12 * mm, bottomMargin=14 * mm, title=title or path.stem,
                          author="Math Inquirers")
    d._foot = footer
    d.build(story, onFirstPage=lambda c, dd: foot(c, d), onLaterPages=lambda c, dd: foot(c, d))


def grid(rows, widths, head=True, extra=None, pad=3):
    t = Table(rows, colWidths=widths, repeatRows=1 if head else 0)
    st = [("GRID", (0, 0), (-1, -1), 0.6, LINE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
          ("TOPPADDING", (0, 0), (-1, -1), pad), ("BOTTOMPADDING", (0, 0), (-1, -1), pad)]
    if head:
        st += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7"))]
    t.setStyle(TableStyle(st + (extra or [])))
    return t


def lesson_title(level, n, t):
    if n == 10 and level <= 4:
        return "Puzzle Day"
    return t


# ---------------------------------------------------------------- 1 parent calendar page
def parent_calendar():
    s = [P("Add your lessons to your calendar", H1),
         P("Math Inquirers November program. One small file puts all 10 lessons in your phone or computer calendar, "
           "with a reminder 1 day before and 1 hour before each lesson. Times are Singapore time.", B),
         P("1. Find your file", H2),
         P("We send you one file named for your child's group, for example <b>L3A_math_inquirers_beta.ics</b> "
           "(L3A = Level 3, Group A). It comes on WhatsApp or by email once your place is confirmed.", B),
         P("2. Open it on your device", H2)]
    rows = [[P("<b>iPhone / iPad</b>", CELL), P("Tap the file, then <b>Add All</b> (or <b>Add to Calendar</b>), and choose your calendar.", CELL)],
            [P("<b>Android (Google Calendar)</b>", CELL), P("Tap the file and choose Google Calendar. If nothing opens: on a computer go to calendar.google.com, "
                                                           "click the gear, <b>Settings</b>, <b>Import and export</b>, <b>Import</b>, pick the file and the calendar.", CELL)],
            [P("<b>Computer (Outlook, Apple Calendar)</b>", CELL), P("Double-click the file. Choose the calendar to add to. Or drag it onto the calendar window.", CELL)]]
    s.append(grid(rows, [48 * mm, 134 * mm], head=False))
    s += [P("3. Check it", H2),
          P("You should see <b>10 events</b>, called <i>Math Inquirers Level N Group A: lesson 1 of 10</i> and so on, each 90 minutes. "
            "Each one lists the Google Meet link, so you can tap it to join (we confirm the link on WhatsApp before lesson 1).", B),
          P("Good to know", H2)]
    s += bullets([
        "<b>Week 5.</b> Nothing is held on Fri 25 Dec (public holiday). Levels 1 to 4 meet <b>Thu 24 Dec</b> and Level 5 meets <b>Wed 23 Dec</b> at the usual time. The file already has the right date.",
        "<b>Opening the file twice</b> does not double your lessons: each lesson has its own code, so a new file replaces the old one. If we change a time we send a new file and tell you on WhatsApp.",
        "<b>Do not delete</b> the reminders: the 1-hour one is the best moment to find the workbook, a pencil, a rubber and a few sheets of blank paper.",
        "Camera is optional and your child may stay muted. Please do not hold paper up to the camera.",
        "Cannot make a lesson? Tell us on WhatsApp " + WA + ". Every lesson is recorded; see the make-up note."])
    s += [Spacer(1, 6), P("<b>Link not working, wrong time, or no file?</b> WhatsApp " + WA + " and we will fix it the same day.", B)]
    build(OUT / "parent_add_to_calendar.pdf", s, footer="Parent guide: add to calendar")


# ---------------------------------------------------------------- 2 make-up note
def makeup():
    s = [P("If your child misses a lesson", H1),
         P("This note matches our terms (section on missed lessons). Nothing here changes them.", BS), Spacer(1, 4),
         P("What happens", H2)]
    s += bullets([
        "<b>Every lesson is recorded.</b> After the lesson we send the private recording link to the families in your child's group on WhatsApp. Your child can watch it and do the lesson's workbook pages and online practice at home.",
        "<b>The same lesson with a later class.</b> If a seat is free in a later class of the <b>same level</b> (for example another group, or a holiday class), your child may attend that same lesson live. We cannot promise a seat: classes are small (up to 8). During the November beta there is one group per level, so for most families the recording is the make-up.",
        "<b>There are no separate make-up lessons</b>, and no refund for a lesson your child does not attend.",
        "<b>School exams.</b> A child who skips lessons during exams can use the recording and, if a seat is free, join the same lesson in another class or in a holiday class."])
    s.append(P("What to do", H2))
    s += bullets([
        "Tell us on WhatsApp " + WA + " <b>before the lesson</b> if you can. A short message is enough: child's name, level and which lesson.",
        "Watch the recording within <b>[X days: owner to confirm]</b> if you can. Recording links stop working 30 days after your group's last lesson, and the recordings are then deleted (see the privacy notice).",
        "Ask us about a seat in a later class on the same WhatsApp number. We will reply with what is free, or tell you plainly that there is none.",
        "If your child is stuck on workbook pages after watching, mark the page, take a photo, and send it to us on WhatsApp. We will give one short written hint. (Please send photos only to this business number.)"])
    s.append(P("What is in a recording", H2))
    s.append(P("The tutor, the shared lesson screen and anyone's voice that is unmuted. Children's cameras are not recorded, and children's names are not shown. "
               "Recordings are shared only with the enrolled families of the same group, through a private link, and never posted publicly. Please do not record or share them yourself.", B))
    s += [Spacer(1, 8), P("Questions? WhatsApp " + WA + ".", B)]
    build(OUT / "makeup_note_parents.pdf", s, footer="Parent note: missed lessons")


# ---------------------------------------------------------------- 3 end-of-program report
SECT = [
    ("What we did", "Topics are listed above. In two or three sentences: the big ideas this child worked on, and one moment worth remembering."),
    ("Strengths we saw", "Two or three specific things, in plain words (for example 'explained why a pattern keeps going', 'kept trying after a wrong turn')."),
    ("A next step", "One thing to grow, framed kindly and doable (for example 'check an answer by working it backwards')."),
    ("Recommended level for next", ""),
    ("Practice ideas for home", "Two or three short ideas for 10 minutes, a few times a week: games, puzzles, everyday maths. No scores needed."),
]
RULE = "No scores, marks, grades or ranking appear on this report, and we do not compare children."


def report_pdf(level):
    ls = LL[str(level)]
    topics = " · ".join(f"{l['n']}. {lesson_title(level, l['n'], l['title'])} ({l['topic']})" for l in ls)
    nxt = [f"Stay at Level {level} (a different set of problems and new term)" if False else f"Continue at Level {level} in the term program",
           f"Move up to Level {level + 1}" if level < 6 else "Continue to the next term at Level 6 (Level 7 is not offered)",
           "Talk to us first (we will call or WhatsApp)"]
    box = lambda: Table([[""]], colWidths=[182 * mm], rowHeights=[24 * mm],
                        style=TableStyle([("BOX", (0, 0), (-1, -1), 0.7, LINE)]))
    s = [P(f"Math Inquirers: Level {level} report", H1),
         grid([[P("<b>Child</b>", CELL), P("________________________", CELL), P("<b>Level / group</b>", CELL), P(f"Level {level}  /  ________", CELL)],
               [P("<b>Dates</b>", CELL), P("________ to ________", CELL), P("<b>Lessons joined live</b>", CELL), P("_____ of 10  (+ recordings: ____)", CELL)],
               [P("<b>Tutor</b>", CELL), P("________________________", CELL), P("<b>Date of report</b>", CELL), P("________", CELL)]],
              [24 * mm, 62 * mm, 38 * mm, 58 * mm], head=False),
         Spacer(1, 3), P("<b>The 10 lessons:</b> " + topics, BS)]
    for h, hint in SECT:
        if h.startswith("Recommended"):
            s.append(P("Recommended level for next", H2))
            s.append(P("  ".join(f"{BOX} {x}" for x in nxt[:2]) + f"  {BOX} {nxt[2]}", CELL))
            s.append(P("Why, in a sentence: ________________________________________________________________________", CELL))
            continue
        s.append(P(h, H2))
        s.append(P(hint, BS))
        s.append(box())
    s += [Spacer(1, 3), P(RULE + " Prepared from the tutor's own observations during the program.", BS)]
    build(OUT / f"end_of_program_report_L{level}.pdf", s, footer=f"End-of-program report, Level {level}", title=f"End-of-program report Level {level}")


def report_docx(level):
    from docx import Document
    from docx.shared import Pt, Mm
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for a in ("left_margin", "right_margin"): setattr(sec, a, Mm(15))
    sec.top_margin = Mm(12); sec.bottom_margin = Mm(12)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
    d.add_heading(f"Math Inquirers: Level {level} report", 1)
    t = d.add_table(rows=3, cols=4); t.style = "Table Grid"
    cells = [["Child", "", "Level / group", f"Level {level} / "], ["Dates", "", "Lessons joined live", "___ of 10 (+ recordings: ___)"],
             ["Tutor", "", "Date of report", ""]]
    for r, row in enumerate(cells):
        for c, v in enumerate(row): t.cell(r, c).text = v
    ls = LL[str(level)]
    p = d.add_paragraph(); p.add_run("The 10 lessons: ").bold = True
    p.add_run("; ".join(f"{l['n']}. {lesson_title(level, l['n'], l['title'])} ({l['topic']})" for l in ls)).font.size = Pt(8.5)
    for h, hint in SECT:
        d.add_heading(h, 2)
        if h.startswith("Recommended"):
            for x in (f"Continue at Level {level} in the term program", f"Move up to Level {level + 1}" if level < 6 else "Continue to the next term at Level 6", "Talk to us first (we will call or WhatsApp)"):
                d.add_paragraph("☐ " + x)
            d.add_paragraph("Why, in a sentence: ")
        else:
            q = d.add_paragraph(hint); q.runs[0].italic = True; q.runs[0].font.size = Pt(8.5)
            d.add_paragraph("\n\n")
    q = d.add_paragraph(RULE + " Prepared from the tutor's own observations during the program."); q.runs[0].font.size = Pt(8.5)
    d.save(OUT / f"end_of_program_report_L{level}.docx")


# ---------------------------------------------------------------- 4 tutor notes sheet (private)
def tutor_notes():
    pg = landscape(A4)
    W = pg[0] - 28 * mm
    s = [P("PRIVATE tutor notes: one child, one program", H1),
         P("<b>Tutor eyes only.</b> Never shown to children or parents; no scores, marks or ranking; keep on paper in the locked folder or in the approved private storage, "
           "never on a personal phone. Use these notes to write the end-of-program report and the next-term group advice, then follow the retention rule in the privacy notice.", BS),
         Spacer(1, 3),
         grid([[P("<b>Child (first name)</b>", CELL), P("", CELL), P("<b>Level / group</b>", CELL), P("", CELL), P("<b>Tutor</b>", CELL), P("", CELL)]],
              [34 * mm, 62 * mm, 28 * mm, 50 * mm, 18 * mm, W - 192 * mm], head=False), Spacer(1, 4)]
    hdr = [P(x, CELLB) for x in ("Lesson", "Live / recording", "Strength seen (specific)", "Stuck on / needed help with", "Engagement and how they worked with others", "Idea for next time")]
    rows = [hdr]
    for n in range(1, 11):
        name = "Lesson %d" % n + ("\n(last lesson)" if n == 10 else "")
        rows.append([P(name.replace("\n", "<br/>"), CELL), P("Live ____<br/>Rec. ____", CELL), "", "", "", ""])
    ws = [24 * mm, 22 * mm, 56 * mm, 56 * mm, 56 * mm, W - 214 * mm]
    t = grid(rows, ws, pad=3)
    t._argH[1:] = [9.5 * mm] * 10
    s += [t, Spacer(1, 5)]
    rows = [[P("<b>What we did, in my words</b>", CELL), P("<b>Strengths to name (2 or 3)</b>", CELL), P("<b>One next step</b>", CELL),
             P("<b>Recommended level</b><br/>%s same level  %s up  %s talk first" % (BOX, BOX, BOX), CELL), P("<b>Practice ideas (2 or 3)</b>", CELL)],
            ["", "", "", "", ""]]
    t2 = grid(rows, [W / 5] * 5)
    s.append(P("End of program (fills the parent report)", H2))
    t2._argH[1] = 22 * mm
    s += [t2, Spacer(1, 3),
          P("Safeguarding or wellbeing concern? Do not write details here: tell the owner the same day by phone or in person.", BS)]
    build(OUT / "tutor_notes_sheet.pdf", s, size=pg, footer="PRIVATE: tutor notes", title="Tutor notes sheet")


# ---------------------------------------------------------------- 5 observation sheet
PARTS = [("Opening", 0, 2), ("Warm-up", 2, 8), ("Sneak-peek", 8, 12), ("Explore", 12, 26), ("Mini-lesson", 26, 36),
         ("Brain break", 36, 40), ("Worksheet", 40, 72), ("Talk time", 72, 79), ("Exit ticket", 79, 85),
         ("Look-back", 85, 88), ("Flex time", 88, 90)]


def observation():
    W = A4[0] - 24 * mm
    s = [P("Tutor trial lesson: observation sheet", H1),
         grid([[P("<b>Level / lesson no.</b>", CELL), P("", CELL), P("<b>Date</b>", CELL), P("", CELL), P("<b>Observer</b>", CELL), P("", CELL)],
               [P("<b>Tutor</b>", CELL), P("", CELL), P("<b>Children (number, ages)</b>", CELL), P("", CELL), P("<b>Platform / device</b>", CELL), P("", CELL)],
               [P("<b>Scheduled start</b>", CELL), P("", CELL), P("<b>Real start (clock)</b>", CELL), P("", CELL), P("<b>Real end (clock)</b>", CELL), P("", CELL)]],
              [30 * mm, 33 * mm, 36 * mm, 30 * mm, 30 * mm, W - 159 * mm], head=False),
         Spacer(1, 3),
         P("Fill in <b>clock times</b> (hh:mm) when each part really starts and ends. The planned column is the Level 1 budget; if the lesson's own budget differs, write it over the top "
           "(budgets are in hp/budget.json and pN/budget.json). Over/under = real minutes minus planned minutes.", BS), Spacer(1, 3)]
    hdr = [P(x, CELLB) for x in ("Part", "Plan min", "Plan from-to", "Real start", "Real end", "Real min", "Over / under", "Stuck points and notes")]
    rows = [hdr]
    for n, a, b in PARTS:
        rows.append([P(n, CELL), P(str(b - a), CELL), P("%d to %d" % (a, b), CELL), "", "", "", "", ""])
    rows.append([P("<b>Total</b>", CELL), P("<b>90</b>", CELL), P("0 to 90", CELL), "", "", "", "", ""])
    ws = [26 * mm, 15 * mm, 21 * mm, 19 * mm, 19 * mm, 15 * mm, 17 * mm, W - 132 * mm]
    t = grid(rows, ws, pad=2)
    t._argH[1:] = [10.5 * mm] * (len(rows) - 1)
    s += [t, Spacer(1, 5)]
    # page 2
    s.append(P("Where children got stuck", H2))
    t = grid([[P(x, CELLB) for x in ("Slide / question", "What happened (who, how many)", "What the tutor did", "Fix needed (wording, time, picture)")]] + [["", "", "", ""]] * 5,
             [30 * mm, 52 * mm, 48 * mm, W - 130 * mm]); t._argH[1:] = [10 * mm] * 5
    s += [KeepTogether([t])]
    t = grid([[P(x, CELLB) for x in ("Question (page / no.)", "Tier (Start here / Think harder / Go further)", "Too hard", "Too easy", "Note")]] + [["", "", BOX, BOX, ""]] * 5,
             [34 * mm, 48 * mm, 16 * mm, 16 * mm, W - 114 * mm]); t._argH[1:] = [9 * mm] * 5
    s += [KeepTogether([P("Questions too hard or too easy", H2), t])]
    s.append(P("Camera, paper and recording rules: seen as planned?", H2))
    rules = ["No child held paper up to the camera (gallery view, tick-against-answer-slide, choral answers, fingers for choices)",
             "'I am still stuck' used: mark on the child's own look-back page, parent photographs it to the business WhatsApp after class",
             "Children could keep camera off and stay muted without comment",
             "Recording on from the start; shows only tutor, shared screen and unmuted voices; the WhatsApp number was not on a recorded slide",
             "Only the lesson window was shared (not the notes, not other tabs); answers did not show early",
             "Private paper: nobody asked to show it; tutor never took a photo or screen capture",
             "No private chat messages with any child; chat set to tutor-visible or off",
             "No personal phone used for photos; Meet link and waiting room worked"]
    rows = [[P("Rule / check", CELLB), P("Yes", CELLB), P("No", CELLB), P("Note", CELLB)]] + [[P(r, CELL), BOX, BOX, ""] for r in rules]
    t = grid(rows, [96 * mm, 11 * mm, 11 * mm, W - 118 * mm], pad=2.5)
    s += [t, P("Technical and tutor-prep notes", H2)]
    t = grid([[P("<b>Sound, screen share, Meet, timer, phone remote</b>", CELL)], [""], [P("<b>Overall: what to change before the real lesson</b>", CELL)], [""]], [W], head=False)
    t._argH[1] = 17 * mm; t._argH[3] = 22 * mm
    s.append(t)
    build(OUT / "tutor_observation_sheet.pdf", s, margin=12, footer="Tutor trial lesson observation sheet", title="Tutor trial lesson observation sheet")


if __name__ == "__main__":
    parent_calendar(); makeup(); tutor_notes(); observation()
    for lv in range(1, 7):
        report_pdf(lv); report_docx(lv)
    print(sorted(p.name for p in OUT.glob("*") if p.is_file()))
