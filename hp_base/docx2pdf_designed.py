import sys, base64, html, os, re
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from playwright.sync_api import sync_playwright

src, out, title, sub, foot = sys.argv[1:6]
landscape = len(sys.argv) > 6 and sys.argv[6] == 'landscape'
doc = Document(src)
E_ = html.escape

def run_html(r):
    parts = []
    for ch in r._r.iterchildren():
        tag = ch.tag.split('}')[1]
        if tag == 't': parts.append(E_(ch.text or ''))
        elif tag == 'tab': parts.append('&emsp;')
        elif tag == 'br' and ch.get(qn('w:type')) != 'page': parts.append('<br>')
        elif tag == 'drawing':
            for b in ch.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
                rid = b.get(qn('r:embed'))
                try:
                    part = doc.part.related_parts[rid]
                    parts.append('<img src="data:%s;base64,%s">' % (part.content_type, base64.b64encode(part.blob).decode()))
                except KeyError: pass
    s = ''.join(parts)
    if not s: return ''
    st = ''
    if r.bold: s = '<b>%s</b>' % s
    if r.italic: s = '<i>%s</i>' % s
    try:
        c = r.font.color.rgb
        if c is not None and str(c) not in ('000000', '1F3864', '1B2A49', '333333'): st += 'color:#%s;' % c
    except Exception: pass
    if r.font.size and r.font.size.pt <= 9.5: st += 'font-size:.86em;'
    return '<span style="%s">%s</span>' % (st, s) if st else s

def has_pb(p):
    return any(br.get(qn('w:type')) == 'page' for br in p._p.iter(qn('w:br')))
def pb_before(p):
    ppr = p._p.pPr
    return ppr is not None and ppr.find(qn('w:pageBreakBefore')) is not None

def para(p, in_cell=False):
    inner = ''.join(run_html(r) for r in p.iter_inner_content() if hasattr(r, '_r')) if hasattr(p, 'iter_inner_content') else ''.join(run_html(r) for r in p.runs)
    text = p.text.strip()
    sn = p.style.name if p.style is not None else 'Normal'
    al = p.alignment
    cls = ''
    if al is not None and int(al) == 1: cls = ' class="c"'
    if not text and '<img' not in inner:
        return '' if in_cell else '<div class="sp"></div>'
    if in_cell:
        return '<div%s>%s</div>' % (cls, inner)
    LV={'core':'Core','think harder':'Think harder','challenge':'Challenge'}
    if text.lower() in LV: return '<div class="lv lv-%s">%s</div>' % (text.lower().replace(' ','-'), LV[text.lower()])
    if re.match(r'^Working:\s*_+\s*$', text): return '<div class="work"><span>Working</span></div>'
    m = re.match(r'^(\d{1,2})\.\s+(.*)$', text, re.S)
    if m and sn == 'Normal' and not p.runs[0].bold:
        rest = re.sub(r'^\s*\d{1,2}\.\s*', '', inner, count=1)
        return '<div class="q"><span class="qn">%s</span><span class="qt">%s</span></div>' % (m.group(1), rest)
    runs = [r for r in p.runs if r.text.strip()]
    allbold = runs and all(r.bold for r in runs)
    mx = max([r.font.size.pt for r in runs if r.font.size] or [0])
    if sn == 'Title': return '<h1 class="title">%s</h1>' % E_(text)
    if sn == 'Heading 1' or (allbold and mx >= 15): return '<h2>%s</h2>' % E_(text)
    if sn == 'Heading 2' or (allbold and mx >= 11.5 and len(text) < 90): return '<h3>%s</h3>' % E_(text)
    if sn.startswith('List'): return '<li>%s</li>' % inner
    ind = p.paragraph_format.left_indent
    st = ' style="padding-left:%dpx"' % min(60, int(ind.pt * 1.1)) if ind and ind.pt > 4 else ''
    if text.startswith('•'):
        return '<li>%s</li>' % re.sub(r'^(<[^>]+>)*•\s*(</[^>]+>)*\s*', '', inner)
    return '<p%s%s>%s</p>' % (cls, st, inner)

def cell_fill(tc):
    tcpr = tc._tc.tcPr
    if tcpr is None: return None
    shd = tcpr.find(qn('w:shd'))
    if shd is None: return None
    f = shd.get(qn('w:fill'))
    return None if not f or f in ('auto', 'FFFFFF') else f

def table(t):
    grid = [int(g.get(qn('w:w')) or 0) for g in t._tbl.tblGrid.iterchildren()] if t._tbl.tblGrid is not None else []
    tot = sum(grid) or 1
    h = ['<table class="tb">']
    if grid: h.append('<colgroup>' + ''.join('<col style="width:%.1f%%">' % (g / tot * 100) for g in grid) + '</colgroup>')
    for ri, row in enumerate(t.rows):
        h.append('<tr>')
        seen = set()
        for tc in row._tr.tc_lst:
            from docx.table import _Cell
            c = _Cell(tc, t)
            span = tc.tcPr.find(qn('w:gridSpan')) if tc.tcPr is not None else None
            cs = ' colspan="%s"' % span.get(qn('w:val')) if span is not None else ''
            f = cell_fill(c)
            cl = ' class="hd"' if f else ''
            body = ''.join(para(p, True) for p in c.paragraphs)
            for nt in c.tables: body += table(nt)
            h.append('<td%s%s>%s</td>' % (cs, cl, body))
        h.append('</tr>')
    h.append('</table>')
    return ''.join(h)

parts = []; inlist = False; first = True
for el in doc.element.body.iterchildren():
    tag = el.tag.split('}')[1]
    if tag == 'p':
        p = Paragraph(el, doc)
        h = para(p)
        if pb_before(p) and not h.startswith('<h3'): parts.append('<div class="brk"></div>')
        if first and p.text.strip() and (p.text.strip().lower() == title.lower() or (p.style.name == 'Title')):
            first = False; h = ''
        elif p.text.strip(): first = False
        if h.startswith('<li>'):
            if not inlist: parts.append('<ul>'); inlist = True
        elif inlist and h != '<div class="sp"></div>':
            parts.append('</ul>'); inlist = False
        parts.append(h)
        if has_pb(p):
            if inlist: parts.append('</ul>'); inlist = False
            parts.append('<div class="brk"></div>')
    elif tag == 'tbl':
        if inlist: parts.append('</ul>'); inlist = False
        parts.append(table(Table(el, doc)))
if inlist: parts.append('</ul>')
# drop page breaks that would leave a nearly empty page
out_parts=[]; since=0
for x in parts:
    if x=='<div class="brk"></div>':
        if since<220: continue
        since=0
    else: since+=len(re.sub('<[^>]+>','',x))
    out_parts.append(x)
parts=out_parts
body = '\n'.join(parts)
body = re.sub(r'(<div class="sp"></div>\s*){2,}', '<div class="sp"></div>', body)
body = re.sub(r'(<div class="brk"></div>\s*)+(<div class="sp"></div>\s*)*$', '', body)
body = re.sub(r'(<div class="brk"></div>\s*)(<div class="sp"></div>\s*)+', r'\1', body)
cover = '<div class="cover"><div class="tag">Math Inquirers Holiday Program</div><h1>%s</h1><div class="sub">%s</div><div class="slog">Notice it. Test it. Explain why.</div></div>' % (E_(title), E_(sub))
css = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), '../fonts/fonts.css')).read() + open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'doc_style.css')).read()
if landscape: css = css.replace('size:A4;', 'size:A4 landscape;')
page = '<!doctype html><html><head><meta charset="utf-8"><style>%s</style></head><body>%s%s</body></html>' % (css, cover, body)
hp = out.replace('.pdf', '.html'); open(hp, 'w').write(page)
ft = '<div style="width:100%%;font-size:8.5px;color:#5B6785;padding:0 14mm;display:flex;justify-content:space-between;font-family:Arial"><span>%s</span><span>Page <span class="pageNumber"></span> of <span class="totalPages"></span></span></div>' % E_(foot)
with sync_playwright() as pw:
    b = pw.chromium.launch(); pg = b.new_page()
    pg.goto('file://' + os.path.abspath(hp), wait_until='networkidle'); pg.wait_for_timeout(800)
    pg.pdf(path=out, format='A4', landscape=landscape, print_background=True, display_header_footer=True, header_template='<div></div>', footer_template=ft, margin={'top': '13mm', 'bottom': '15mm', 'left': '13mm', 'right': '13mm'})
    b.close()
print('ok', out)
