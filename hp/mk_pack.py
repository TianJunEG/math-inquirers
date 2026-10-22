import json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
NAVY=RGBColor(0x1B,0x2A,0x49)
def base(d):
    s=d.styles['Normal'];s.font.name='Calibri';s.font.size=Pt(11)
    for sec in d.sections: sec.left_margin=sec.right_margin=Cm(2);sec.top_margin=sec.bottom_margin=Cm(1.8)
def H(d,t,l=1):
    h=d.add_heading(t,l)
    for r in h.runs: r.font.color.rgb=NAVY;r.font.name='Calibri'
    return h
def P(d,t,b=False,i=False,sz=None):
    p=d.add_paragraph();r=p.add_run(t);r.bold=b;r.italic=i
    if sz:r.font.size=Pt(sz)
    return p
def shade(c,col):
    tc=c._tc.get_or_add_tcPr();s=OxmlElement('w:shd');s.set(qn('w:val'),'clear');s.set(qn('w:fill'),col);tc.append(s)
def table(d,rows,widths=None,head=True):
    t=d.add_table(rows=len(rows),cols=len(rows[0]));t.style='Table Grid';t.autofit=False
    for i,r in enumerate(rows):
        for j,v in enumerate(r):
            c=t.cell(i,j);c.text=''
            run=c.paragraphs[0].add_run(str(v))
            if head and i==0: run.bold=True;shade(c,'FFC93C')
            if widths: c.width=Cm(widths[j])
    if widths:
        for j,wd in enumerate(widths): t.columns[j].width=Cm(wd)
    for r in t.rows:
        for c in r.cells:
            for pp in c.paragraphs:
                for rr in pp.runs: rr.font.size=Pt(10)
    return t

# ---------- EMAIL PACK ----------
W=[
 (1,"Lessons 1 and 2","Number Patterns and Number Bonds",
  "Count with the Frog; Number Bond Inquiry",
  "Cut-out kit pages (ten-frames, counters, number strips) are at the end of the week 1 PDF. Please cut them out and keep them in one envelope; we use them every week.",
  "Ask your child: 'Pick a number from 1 to 10. How many ways can you split it into two parts? How do you know you found them all?'",
  ["3bJhBJ37ZoAfsuDo4LgWeH","1fsJ5Fj2abUf27LYUmuEqN"]),
 (2,"Lessons 3 and 4","Picture Graphs and Repeating and Growing Patterns",
  "Fruit Party Graphs; Pattern Inquiry",
  "No new cut-outs this week. Please keep the week 1 envelope handy (counters are used again).",
  "Ask your child: 'Make a pattern with spoons and forks. What will the 10th item be? How do you know without drawing it all?'",
  ["4BN9fWMEnxLFD1bPLiF8Yf","McjqWjXihwfoCzT9BNfr7W"]),
 (3,"Lessons 5 and 6","Working Backwards and Money",
  "Mystery Number; Coin Combos",
  "The play coins and the play notes ($1 coins and $2, $5, $10 and $50 notes) are on the last two pages of the week 3 PDF. Please cut them out and add them to the envelope.",
  "Ask your child: 'I am thinking of a number. I add 5 and get 12. What is my number? How can you check it?' Then play shop: give a price like $17 or 65¢ and ask your child to pay with the fewest coins and notes.",
  ["HqYXBEwzKv13Ne8YX9BfZS","6epfeWSo1ECFc6Dp7Xetc5"]),
 (4,"Lessons 7 and 8","2D Shapes and Logic and Sorting",
  "Shape Safari; Logic and Sorting",
  "The sorting tiles are in the week 1 kit pages. Please bring the envelope to class or keep it by your child for online lessons.",
  "Ask your child: 'Find three things at home. Which one is the odd one out? There can be more than one right answer. Explain why.'",
  ["RL9jDmrZwdbjVnL98L2tEp","UjMh35Wk2hZfT59b3W1TtV"]),
 (5,"Lessons 9 and 10","Length and Challenge Day (mixed puzzles)",
  "Measure It; Challenge Day",
  "The paper cube strips are on the cut-outs page (included again in the week 5 PDF). The Certificate of Completion is the last page.",
  "Ask your child: 'Measure the table with your hand span, then with a pencil. Why are the two numbers different?'",
  ["V7cAm5E5H9i2N8bp3ynmsi","KCqynhP5sQ5BiYCeUzgJNs"]),
]
d=Document();base(d)
H(d,"Math Inquirers Holiday Program: Weekly Email Pack",0)
P(d,"Notice it. Test it. Explain why.",i=True)
H(d,"How to use this pack",1)
for x in [
 "Send the parent email and the Week N Student Pages PDF the day before the first lesson of that week.",
 "Send the tutor email with the Week N Tutor Scripts and Answers PDF at least two days before. Tutors only; the answers must not go to parents.",
 "Put parents in BCC so no one sees another family's email address. Get parents' consent to email before you start (PDPA).",
 "Replace the words in [square brackets] before sending.",
 "The PDFs are in the weekly folder: Week1_Student_Pages.pdf to Week5_Student_Pages.pdf and Week1_Tutor_Scripts_and_Answers.pdf to Week5_Tutor_Scripts_and_Answers.pdf.",
 "The slide deck links in the tutor emails are private to you. Share each deck with the tutor first, or send them the file.",
]: d.add_paragraph(x,style='List Bullet')
for w,ls,top,names,kit,prompt,ids in W:
    d.add_page_break()
    H(d,f"Week {w}: {ls}",1)
    H(d,"Parent email",2)
    P(d,f"Subject: Math Inquirers Week {w}: {top}",b=True)
    P(d,f"Attachment: Week{w}_Student_Pages.pdf",i=True)
    body=f"""Dear parents,

This week our Math Inquirers explore {top.lower()} in {ls.lower()}.

Your child's workbook pages for this week are attached. Please print them single-sided if you can. Children will work on the pages in class, so there is no need to finish them in advance.

{kit}

One thing to try at home (5 minutes, no pencil needed): {prompt[0].lower()+prompt[1:]}

We are not looking for speed. We are looking for a child who can say what they noticed, test it, and explain why. If your child is stuck, please ask "What do you notice?" before you give the answer.

Warm regards,
[Your name]
Preciouslyours Pte Ltd
[Contact number]"""
    for para in body.split("\n\n"): P(d,para)
    H(d,"Tutor email",2)
    P(d,f"Subject: Math Inquirers Week {w} tutor pack: {ls}",b=True)
    P(d,f"Attachments: Week{w}_Tutor_Scripts_and_Answers.pdf, Week{w}_Student_Pages.pdf",i=True)
    tb=f"""Hi [tutor name],

Here is your pack for week {w}: {names}.

Attached are the minute-by-minute scripts with answers for both lessons, and the student pages. The slide decks are linked below.
  Lesson {2*w-1} deck: https://claude.ai/artifact/{ids[0]}
  Lesson {2*w} deck: https://claude.ai/artifact/{ids[1]}

Before the lesson, please:
  1. Read both scripts once and try the Go further questions yourself.
  2. Check you have counters, coins and cut-outs ready (see the materials list in the guide).
  3. After each lesson, fill in the Progress Notes for each child (30 seconds each) and send me the sheet.

Please do not forward the answer pages to parents.

Thank you,
[Your name]"""
    for para in tb.split("\n\n"): P(d,para)
d.add_page_break()
H(d,"Week 5 closing email to parents",1)
P(d,"Subject: Math Inquirers: thank you, and your child's certificate",b=True)
for para in ["Dear parents,","Thank you for joining us for the Math Inquirers Holiday Program. Your child's Certificate of Completion is the last page of the Week 5 PDF. Please print it and add your child's name.","We would love your feedback in two minutes: [feedback form link]. What did your child enjoy most? What should we change?","If your child showed a love for puzzles, we may invite them to a short challenge session in the next holiday. We will contact you if so.","Warm regards,\n[Your name]"]: P(d,para)
d.save("Math_Inquirers_Level1_Weekly_Email_Pack.docx")

# ---------- PROGRESS NOTES ----------
d=Document();base(d)
sec=d.sections[0];sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width,sec.page_height=sec.page_height,sec.page_width
H(d,"Math Inquirers: Progress Notes",0)
P(d,"A short note after each lesson on what each child tried, enjoyed and found tricky. It feeds the weekly message to parents. It is not a test and it does not rank or label anyone. Use usernames or first names only. Keep it to 30 seconds per child.")
H(d,"Four things worth noting",1)
table(d,[["Code","What to look for","Example"],
 ["A","Had a go","Tried the Start here questions, used counters, a drawing or words."],
 ["B","Explained a step","Said how they got an answer, or why a pattern works."],
 ["C","Tried something harder","Chose Think harder, Go further or the Investigation, in their own time."],
 ["D","Needs another look","Something to revisit next lesson, with a gentle way to help."]],[2,7,14])
P(d,"")
H(d,"Lesson log: copy one page per lesson",1)
P(d,"Lesson number: ______   Date: ____________   Tutor: ____________________")
rows=[["Child (username)","A","B","C","D","What they said or did, and something to celebrate"]]+[[f"{i}.","","","","",""] for i in range(1,9)]
table(d,rows,[5,1.5,1.5,1.5,1.5,13])
P(d,"")
d.add_page_break()
H(d,"Programme summary: one row per child",1)
P(d,"Tick or jot one line per lesson. At the end of the programme this becomes the closing note to the family.")
rows=[["Child"]+[f"L{i}" for i in range(1,11)]+["Favourite moment","Next step to suggest"]]+[[f"{i}."]+[""]*10+["",""] for i in range(1,9)]
table(d,rows,[4]+[1.4]*10+[3,3])
P(d,"")
H(d,"How to use the notes",1)
for x in ["Write what you saw, not a score. Quote the child's own words where you can.",
"Always start the weekly message with something the child did well.",
"If a child tried something harder, tell the family so they can enjoy it with them, and mention the next level only if the family asks.",
"If a child found a lesson tricky, say what helped and what to try at home. Do not compare children or share notes between families.",
"Notes are private to the tutor and the child's own family. Do not record anything beyond learning in class."]: d.add_paragraph(x,style='List Bullet')
d.save("Math_Inquirers_Progress_Notes.docx")
